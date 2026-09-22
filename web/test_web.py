"""
Automated Verification Test Suite for Comment Absorber Web
Tests Flask endpoints, simulation streaming, KPI stats, and Excel/CSV exports.
"""

import sys
import time
import unittest
from pathlib import Path
import openpyxl

from app_web import app, session

class TestCommentAbsorberWeb(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.client = app.test_client()
        session.clear()

    def tearDown(self):
        session.clear()

    def test_01_index_page(self):
        """Verify the main web dashboard loads with 200 OK and expected elements."""
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        html = resp.data.decode("utf-8")
        self.assertIn("Comment Absorber", html)
        self.assertIn("CURRENTLY COLLECTED COMMENT", html)
        self.assertIn("Comments Feed", html)

    def test_02_start_simulator_collection(self):
        """Verify starting collection in Simulator mode streams comments."""
        resp = self.client.post("/api/collect/start", json={
            "url": "https://www.facebook.com/test/posts/12345",
            "mode": "simulator",
            "max_comments": 25,
            "speed": 0.05
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertTrue(data["success"])

        # Wait for some comments to be simulated
        time.sleep(0.8)

        state_resp = self.client.get("/api/collect/state")
        self.assertEqual(state_resp.status_code, 200)
        state_data = state_resp.get_json()
        
        self.assertGreater(state_data["stats"]["count"], 0)
        self.assertIsNotNone(state_data["stats"]["latest_comment"])
        self.assertIn("user_name", state_data["stats"]["latest_comment"])
        self.assertIn("message", state_data["stats"]["latest_comment"])
        print(f"\n[+] Verified: {state_data['stats']['count']} comments streamed in real-time.")
        print(f"[+] Latest Comment in Spotlight: {state_data['stats']['latest_comment']['user_name']}: {state_data['stats']['latest_comment']['message'][:40]}...")

    def test_03_export_excel_and_csv(self):
        """Verify that collected comments can be exported to Excel and CSV."""
        # Add sample comments
        from web_collector import WebComment
        session.on_new_comment(WebComment(
            index=1,
            comment_id="test_1",
            user_name="Juan Dela Cruz",
            message="Interested! Magkano po shipping? ✨",
            created_time="2026-09-22 10:30:00",
            timestamp_raw=1790073000.0
        ))
        session.on_new_comment(WebComment(
            index=2,
            comment_id="test_2",
            user_name="Maria Santos",
            message="Mine 1 set please. COD available? 📦",
            created_time="2026-09-22 10:35:00",
            timestamp_raw=1790073300.0
        ))

        # 1. Test Excel export
        excel_resp = self.client.get("/api/export/excel?sort=oldest")
        self.assertEqual(excel_resp.status_code, 200)
        self.assertEqual(excel_resp.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
        # Save temp and inspect with openpyxl
        temp_excel = Path("test_output.xlsx")
        temp_excel.write_bytes(excel_resp.data)
        
        wb = openpyxl.load_workbook(temp_excel)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "User")
        self.assertEqual(ws.cell(row=1, column=2).value, "Comment")
        self.assertEqual(ws.cell(row=1, column=3).value, "Date")
        self.assertEqual(ws.cell(row=2, column=1).value, "Juan Dela Cruz")
        self.assertIn("Interested", ws.cell(row=2, column=2).value)
        wb.close()
        temp_excel.unlink(missing_ok=True)
        print("[+] Verified Excel export: Proper headers ('User', 'Comment', 'Date') & rows match.")

        # 2. Test CSV export
        csv_resp = self.client.get("/api/export/csv?sort=oldest")
        self.assertEqual(csv_resp.status_code, 200)
        csv_text = csv_resp.data.decode("utf-8-sig")
        self.assertIn('"User","Comment","Date"', csv_text)
        self.assertIn("Juan Dela Cruz", csv_text)
        self.assertIn("Maria Santos", csv_text)
        print("[+] Verified CSV export: CSV structure matches specification.")

    def test_04_stop_and_clear(self):
        """Verify stop and clear actions reset state."""
        self.client.post("/api/collect/start", json={"mode": "simulator", "max_comments": 10, "speed": 0.5})
        time.sleep(0.3)
        stop_resp = self.client.post("/api/collect/stop")
        self.assertTrue(stop_resp.get_json()["success"])
        
        clear_resp = self.client.post("/api/collect/clear")
        self.assertTrue(clear_resp.get_json()["success"])
        
        state = self.client.get("/api/collect/state").get_json()
        self.assertEqual(state["stats"]["count"], 0)
        self.assertEqual(state["stats"]["status"], "IDLE")
    def test_05_theme_and_cors_support(self):
        """Verify theme toggle button, GitHub Pages elements, and CORS headers."""
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        html = resp.data.decode("utf-8")
        self.assertIn("theme-toggle-btn", html)
        self.assertIn("data-theme=\"light\"", html)
        self.assertIn("gh-pages-notice", html)
        self.assertIn("gh-modal", html)
        
        # Verify CORS headers
        status_resp = self.client.get("/api/auth/status")
        self.assertEqual(status_resp.headers.get("Access-Control-Allow-Origin"), "*")
        print("\n[+] Verified: Theme toggle, GitHub Pages elements, and CORS headers are active.")

if __name__ == "__main__":
    unittest.main()
