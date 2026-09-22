"""
Unit tests for ExcelExporter.
Verifies openpyxl output structure, strictly 3 columns (User, Comment, Date),
oldest-to-newest sorting, freeze panes, auto-filters, text wrapping,
and preservation of emojis, Filipino text, and Unicode.
"""

from datetime import datetime, timezone
from pathlib import Path
import openpyxl
import pytest

from app.excel_exporter import ExcelExporter
from app.models import Comment


class TestExcelExporter:

    def test_export_creates_valid_file_with_exact_three_columns(self, tmp_path):
        c1 = Comment(
            comment_id="c_1",
            user_name="Juan Dela Cruz",
            message="Magandang araw po sa inyong lahat! Mabuhay! 🇵🇭 🎉",
            created_time=datetime(2026, 9, 12, 8, 0, 0, tzinfo=timezone.utc)
        )
        c2 = Comment(
            comment_id="c_2",
            user_name="Sarah Connor",
            message="Line one.\nLine two with emoji 🚀\nLine three.",
            created_time=datetime(2026, 9, 12, 8, 15, 0, tzinfo=timezone.utc)
        )
        c3 = Comment(
            comment_id="c_3",
            user_name="Facebook User",  # Privacy fallback
            message="Testing illegal char:\x00\x08cleaned!",
            created_time=datetime(2026, 9, 12, 7, 30, 0, tzinfo=timezone.utc)  # Oldest!
        )

        exporter = ExcelExporter(output_dir=tmp_path)
        output_file = exporter.export([c1, c2, c3], target_path=tmp_path / "test_out.xlsx")

        assert output_file.exists()

        # Read workbook with openpyxl to inspect structure
        wb = openpyxl.load_workbook(output_file)
        ws = wb.active
        assert ws.title == "Comments"

        # Check Header
        assert ws.max_column == 3
        headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
        assert headers == ["User", "Comment", "Date"]

        # Check Header Styling
        header_cell = ws.cell(row=1, column=1)
        assert header_cell.font.bold is True
        assert header_cell.font.color.rgb == "00FFFFFF" or header_cell.font.color.rgb == "FFFFFF"
        assert header_cell.fill.start_color.rgb == "001B365D" or header_cell.fill.start_color.rgb == "1B365D"

        # Check Freeze Panes
        assert ws.freeze_panes == "A2"

        # Check AutoFilter
        assert ws.auto_filter.ref is not None
        assert "A1:C" in ws.auto_filter.ref

        # Check Total Rows (1 header + 3 data rows = 4)
        assert ws.max_row == 4

        # Check Oldest to Newest Sorting
        # Row 2 should be c3 (07:30:00)
        assert ws.cell(row=2, column=1).value == "Facebook User"
        assert ws.cell(row=2, column=2).value == "Testing illegal char:cleaned!"
        assert ws.cell(row=2, column=3).value == "2026-09-12 07:30:00"

        # Row 3 should be c1 (08:00:00)
        assert ws.cell(row=3, column=1).value == "Juan Dela Cruz"
        assert "Magandang araw po" in ws.cell(row=3, column=2).value
        assert "🇵🇭" in ws.cell(row=3, column=2).value
        assert ws.cell(row=3, column=3).value == "2026-09-12 08:00:00"

        # Row 4 should be c2 (08:15:00)
        assert ws.cell(row=4, column=1).value == "Sarah Connor"
        assert "\n" in ws.cell(row=4, column=2).value
        assert "🚀" in ws.cell(row=4, column=2).value
        assert ws.cell(row=4, column=3).value == "2026-09-12 08:15:00"

        # Check Comment Column Wrap Text & Width
        comment_cell = ws.cell(row=3, column=2)
        assert comment_cell.alignment.wrap_text is True
        assert ws.column_dimensions["B"].width >= 70

    def test_filename_pattern(self, tmp_path):
        exporter = ExcelExporter(output_dir=tmp_path)
        dt = datetime(2026, 9, 12, 19, 55, 0)
        fname = exporter.generate_filename(dt)
        assert fname == "Facebook_Comments_2026-09-12_19-55-00.xlsx"

    def test_filename_pattern_with_author_name(self, tmp_path):
        exporter = ExcelExporter(output_dir=tmp_path)
        dt = datetime(2026, 9, 12, 19, 55, 0)
        
        # Test 1: Author with apostrophe and spaces
        fname1 = exporter.generate_filename(dt, author_name="Alisha's Gadgets Cellphone Repair")
        assert fname1 == "Alishas_Gadgets_Cellphone_Repair_Comments_2026-09-12_19-55-00.xlsx"

        # Test 2: Standard personal name
        fname2 = exporter.generate_filename(dt, author_name="Jay Paul Torres Benolirao")
        assert fname2 == "Jay_Paul_Torres_Benolirao_Comments_2026-09-12_19-55-00.xlsx"

        # Test 3: Author with illegal file system characters
        fname3 = exporter.generate_filename(dt, author_name='Page\\Name/With:Invalid*Chars?"<>|')
        assert fname3 == "PageNameWithInvalidChars_Comments_2026-09-12_19-55-00.xlsx"

        # Test 4: Blank or None author falls back to Facebook_Comments_...
        fname4 = exporter.generate_filename(dt, author_name="   ")
        assert fname4 == "Facebook_Comments_2026-09-12_19-55-00.xlsx"

