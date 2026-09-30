"""
Tests for Exact Original Comment Preservation, Language Integrity,
Developer Diagnostics, and Excel/CSV Export Fidelity.
"""

from datetime import datetime, timezone
from pathlib import Path
import openpyxl
import pytest

from app.models import Comment
from web.web_collector import WebComment, ExcelReportExporter
from web.app_web import CollectionSession


class TestOriginalCommentPreservation:

    def test_original_text_strictly_preferred_over_translated_text(self):
        """
        Verifies that when both original_text and translated_text exist,
        original_text is strictly preserved as the primary message for display,
        storage, and export.
        """
        cebuano_original = "Grabe ka gwapa ani uy 😍"
        english_translation = "You are so beautiful 😍"

        comment = WebComment(
            index=1,
            comment_id="c_101",
            user_name="Juan Dela Cruz",
            message=english_translation,  # simulated initial visual field
            created_time="2026-09-30 20:00:00",
            timestamp_raw=1790000000.0,
            original_text=cebuano_original,
            translated_text=english_translation,
            is_translation=True,
            original_field_used="original_text",
            raw_source_text=cebuano_original
        )

        # Primary message MUST be the exact original comment text
        assert comment.message == cebuano_original
        assert comment.original_text == cebuano_original
        assert comment.is_translation is True

        # Diagnostic payload must contain all required developer inspection fields
        diag = comment.debug_diagnostic
        assert diag["comment_id"] == "c_101"
        assert diag["raw_source_text"] == cebuano_original
        assert diag["selected_display_text"] == cebuano_original
        assert diag["translation_detected"] == "YES"
        assert diag["original_field_used"] == "original_text"

    def test_multilingual_comments_verbatim_preservation(self):
        """
        Verifies preservation of exact spelling, capitalization, punctuation, emojis,
        slang, repeated letters, and line breaks in various languages (Bisaya, Tagalog, Japanese, Spanish).
        """
        test_samples = [
            ("Cebuano", "Maayo kaayo ni bai hahaha 😂😂😂 grabe jud ka"),
            ("Tagalog", "Sobrang ganda naman nito sis!! Pwede pa-mine?? 💕"),
            ("Japanese", "とても綺麗ですね！🌸✨ 素晴らしい写真をありがとう。"),
            ("Spanish", "¡¡Muchísimas felicidades amigo!! 🎉🎊 Nos vemos pronto."),
            ("Punctuation & Slang", "HAHAHAHAHA OMG!!??? 😭😭 legit jud sha 100%!!!\nSecond line here.")
        ]

        for idx, (lang, raw_text) in enumerate(test_samples, start=1):
            comment = WebComment(
                index=idx,
                comment_id=f"c_{idx}",
                user_name=f"User {lang}",
                message=raw_text,
                created_time="2026-09-30 20:05:00",
                timestamp_raw=1790000000.0 + idx,
                original_text=raw_text,
                is_translation=False,
                original_field_used="raw_source"
            )

            # Strict character-for-character equality
            assert comment.message == raw_text
            assert comment.original_text == raw_text
            assert comment.debug_diagnostic["translation_detected"] == "NO"

    def test_excel_export_preserves_original_comments(self, tmp_path):
        """
        Verifies that Excel exports output the exact original text,
        even when simulated translation fields are present.
        """
        c1 = WebComment(
            index=1,
            comment_id="c_201",
            user_name="Juan",
            message="This is very good bro hahaha",  # translated
            created_time="2026-09-30 19:30:00",
            timestamp_raw=1790000000.0,
            original_text="Maayo kaayo ni bai hahaha 😍",
            translated_text="This is very good bro hahaha",
            is_translation=True,
            original_field_used="original_text"
        )
        c2 = WebComment(
            index=2,
            comment_id="c_202",
            user_name="Maria",
            message="You are so beautiful 😍",
            created_time="2026-09-30 19:35:00",
            timestamp_raw=1790000005.0,
            original_text="Grabe ka gwapa ani uy 😍",
            translated_text="You are so beautiful 😍",
            is_translation=True,
            original_field_used="original_text"
        )

        out_file = tmp_path / "original_comments_test.xlsx"
        ExcelReportExporter.export([c1, c2], out_file, sort_order="oldest", include_names=True)

        assert out_file.exists()

        wb = openpyxl.load_workbook(out_file)
        ws = wb.active
        assert ws.title in ("Facebook Comments", "Comments")

        # Check Row 2 (Juan's comment)
        assert ws.cell(row=2, column=1).value == "Juan"
        assert ws.cell(row=2, column=2).value == "Maayo kaayo ni bai hahaha 😍"

        # Check Row 3 (Maria's comment)
        assert ws.cell(row=3, column=1).value == "Maria"
        assert ws.cell(row=3, column=2).value == "Grabe ka gwapa ani uy 😍"

    def test_collection_session_tracks_top_level_and_replies_separately(self):
        """
        Verifies that CollectionSession tracks top-level comments and replies
        separately and deduplicates strictly by comment ID.
        """
        sess = CollectionSession()

        top_comment = WebComment(
            index=1,
            comment_id="top_1",
            user_name="Parent Commenter",
            message="Kinsay moadto unya? 👀",
            created_time="2026-09-30 18:00:00",
            timestamp_raw=1790000000.0,
            original_text="Kinsay moadto unya? 👀",
            is_reply=False
        )

        reply_comment = WebComment(
            index=2,
            comment_id="reply_1",
            user_name="Child Commenter",
            message="Ako bai! Sabay ta hahahaha",
            created_time="2026-09-30 18:05:00",
            timestamp_raw=1790000001.0,
            original_text="Ako bai! Sabay ta hahahaha",
            is_reply=True,
            parent_id="top_1"
        )

        # Identical comment text from a DIFFERENT author
        duplicate_text_different_user = WebComment(
            index=3,
            comment_id="reply_2",
            user_name="Another User",
            message="Ako bai! Sabay ta hahahaha",
            created_time="2026-09-30 18:06:00",
            timestamp_raw=1790000002.0,
            original_text="Ako bai! Sabay ta hahahaha",
            is_reply=True,
            parent_id="top_1"
        )

        sess.on_new_comment(top_comment)
        sess.on_new_comment(reply_comment)
        sess.on_new_comment(duplicate_text_different_user)

        stats = sess.get_stats_dict()
        assert stats["count"] == 3
        assert stats["top_level_count"] == 1
        assert stats["replies_count"] == 2
        assert stats["unique_authors"] == 3
