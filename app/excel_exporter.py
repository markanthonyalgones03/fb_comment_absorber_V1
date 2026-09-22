"""
Excel Exporter for Facebook Comment Collector using openpyxl.
Generates clean, professionally formatted .xlsx spreadsheets
with exactly three columns: User, Comment, Date, sorted oldest to newest.
"""

from datetime import datetime, timezone
import re
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Comment
from app.utils import format_excel_date, sanitize_for_excel, sort_comments_oldest_first


class ExcelExporter:
    """
    Handles generation and formatting of the comments Excel spreadsheet.
    """

    def __init__(self, output_dir: Optional[Path] = None):
        if output_dir is None:
            output_dir = Path("output")
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_filename(
        self,
        timestamp: Optional[datetime] = None,
        author_name: Optional[str] = None
    ) -> str:
        """
        Generates filename according to pattern:
        {Author_Name}_Comments_YYYY-MM-DD_HH-MM-SS.xlsx
        or fallback: Facebook_Comments_YYYY-MM-DD_HH-MM-SS.xlsx
        """
        if timestamp is None:
            timestamp = datetime.now()
        ts_str = timestamp.strftime("%Y-%m-%d_%H-%M-%S")
        if author_name and author_name.strip():
            clean_name = re.sub(r'[\\/*?:"<>|\'’]', '', author_name.strip())
            clean_name = re.sub(r'\s+', '_', clean_name).strip('_')
            if clean_name:
                return f"{clean_name}_Comments_{ts_str}.xlsx"
        return f"Facebook_Comments_{ts_str}.xlsx"

    def export(
        self,
        comments: List[Comment],
        target_path: Optional[Path] = None,
        author_name: Optional[str] = None
    ) -> Path:
        """
        Exports comments to an Excel file with required formatting.
        Guarantees:
        1. Duplicate comments are already removed.
        2. Sorted strictly OLDEST -> NEWEST.
        3. Exactly three columns: User, Comment, Date.
        4. Frozen header row, auto-filter enabled, wrapped comments, safe for emojis/Filipino text.
        """
        # Ensure Oldest -> Newest sorting
        sorted_comments = sort_comments_oldest_first(comments)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comments"

        # Ensure gridlines are visible
        ws.views.sheetView[0].showGridLines = True

        # Header Definition
        headers = ["User", "Comment", "Date"]
        ws.append(headers)

        # Style Header Row
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        header_border = Border(
            bottom=Side(style="medium", color="0D1E36")
        )

        for col_num in range(1, 4):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = header_border

        ws.row_dimensions[1].height = 28

        # Styling for data rows
        data_font = Font(name="Segoe UI", size=10)
        user_alignment = Alignment(horizontal="left", vertical="top")
        comment_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        date_alignment = Alignment(horizontal="center", vertical="top")
        thin_border = Border(
            bottom=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB")
        )

        # Populate Data Rows
        row_idx = 2
        for comment in sorted_comments:
            user_text = sanitize_for_excel(comment.user_name)
            comment_text = sanitize_for_excel(comment.message)
            date_text = format_excel_date(comment.created_time)

            row_data = [user_text, comment_text, date_text]
            ws.append(row_data)

            # Apply cell styles
            c_user = ws.cell(row=row_idx, column=1)
            c_user.font = data_font
            c_user.alignment = user_alignment
            c_user.border = thin_border

            c_comment = ws.cell(row=row_idx, column=2)
            c_comment.font = data_font
            c_comment.alignment = comment_alignment
            c_comment.border = thin_border

            c_date = ws.cell(row=row_idx, column=3)
            c_date.font = data_font
            c_date.alignment = date_alignment
            c_date.border = thin_border

            row_idx += 1

        total_rows = len(sorted_comments) + 1  # including header

        # Column Widths
        # User: 26, Comment: 75 (wide for comfortable reading), Date: 22
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 75
        ws.column_dimensions["C"].width = 22

        # Freeze Header Row
        ws.freeze_panes = "A2"

        # Enable AutoFilter on A1:C{total_rows}
        ws.auto_filter.ref = f"A1:C{max(total_rows, 2)}"

        # Determine Destination Path
        if target_path is None:
            filename = self.generate_filename(author_name=author_name)
            destination = self.output_dir / filename
        else:
            destination = Path(target_path)
            destination.parent.mkdir(parents=True, exist_ok=True)

        wb.save(destination)
        return destination.resolve()
